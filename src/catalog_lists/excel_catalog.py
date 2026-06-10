"""
Import a 名录 选型对照表 **.xlsx** as model-level products + a CatalogList.

Unlike the PDF importer (which only extracts codes and matches them to
already-crawled products), this Excel has the model details inline, so we
*create* the Product rows (tagged ``granularity="model"``) AND register them
in a named CatalogList — that's the 名录→型号 scope.

Reads the flat ``总览`` sheet (类别 / 型号定位 / 型号 / 核心定位) and enriches
each model's description with scenarios from ``选型建议`` when present.
"""

from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from sqlalchemy import delete, select

from ..storage import get_db
from ..storage.models import CatalogEntry, CatalogList, Product

logger = logging.getLogger(__name__)


@dataclass
class CatalogXlsxReport:
    catalog_name: str
    source_file: str
    models_total: int = 0
    by_category: dict[str, int] = field(default_factory=dict)


def _clean(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\r", " ").replace("\n", " ").strip()


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _code_token(model: str) -> str:
    """'UNIS S5800-56T-HI-G' -> 'S5800-56T-HI-G' for loose scenario matching."""
    t = model
    for noise in ("UNIS", "Server", "Storage"):
        t = t.replace(noise, " ")
    parts = [p for p in t.split() if p]
    return parts[0] if parts else ""


def _read_scenarios(wb) -> dict[str, list[str]]:
    """选型建议 sheet -> {推荐型号串: [场景说明 lines]} (best-effort, optional)."""
    out: dict[str, list[str]] = {}
    if "选型建议" not in wb.sheetnames:
        return out
    rows = list(wb["选型建议"].iter_rows(values_only=True))
    for raw in rows[2:]:
        cells = [_clean(c) for c in raw]
        scen = cells[0] if len(cells) > 0 else ""
        models = cells[1] if len(cells) > 1 else ""
        note = cells[2] if len(cells) > 2 else ""
        if not models:
            continue
        out.setdefault(models, []).append(scen + (("：" + note) if note else ""))
    return out


def import_catalog_xlsx(
    xlsx_path: str | Path,
    *,
    name: str,
    notes: str | None = None,
    replace: bool = True,
) -> CatalogXlsxReport:
    """Import an Excel 名录 as model-level Products + a named CatalogList."""
    import openpyxl  # local import

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if "总览" not in wb.sheetnames:
        raise ValueError("名录 xlsx 需要一个『总览』sheet 列出型号;未找到。")

    rows = list(wb["总览"].iter_rows(values_only=True))
    headers = [_clean(h) for h in rows[1]] if len(rows) > 1 else []

    def find_col(keywords: list[str], default: int) -> int:
        # Exact header match first, so "型号" can't be hijacked by "型号定位".
        for i, h in enumerate(headers):
            if h in keywords:
                return i
        for i, h in enumerate(headers):
            if any(k in h for k in keywords):
                return i
        return default

    c_cat = find_col(["类别"], 0)
    c_pos = find_col(["型号定位", "定位"], 2)
    c_model = find_col(["型号"], 3)
    c_core = find_col(["核心定位", "核心", "说明"], 4)

    scen_map = _read_scenarios(wb)

    def scenarios_for(model: str) -> list[str]:
        tok = _code_token(model)
        if not tok:
            return []
        bits: list[str] = []
        for key, lines in scen_map.items():
            if tok in key:
                bits.extend(lines)
        return bits

    # ---- parse 总览 (forward-fill merged 类别) ------------------------------
    parsed: list[tuple[str, str, str, str]] = []
    last_cat = ""
    for raw in rows[2:]:
        cells = list(raw)

        def g(i: int) -> str:
            return _clean(cells[i]) if i < len(cells) else ""

        cat = g(c_cat) or last_cat
        last_cat = cat
        model = g(c_model)
        if not model:
            continue
        parsed.append((cat, model, g(c_pos), g(c_core)))

    db = get_db()
    report = CatalogXlsxReport(catalog_name=name, source_file=str(xlsx_path))

    with db.session() as s:
        existing = s.scalar(select(CatalogList).where(CatalogList.name == name))
        if existing and not replace:
            raise ValueError(f"名录 {name!r} 已存在;replace=True 才能覆盖。")
        if existing:
            s.execute(delete(CatalogEntry).where(CatalogEntry.catalog_id == existing.id))
            existing.source_file = str(xlsx_path)
            existing.source_sha256 = _sha256(xlsx_path)
            existing.extractor = "excel"
            existing.imported_at = datetime.utcnow()
            existing.notes = notes
            catalog = existing
        else:
            catalog = CatalogList(
                name=name, source_file=str(xlsx_path), source_sha256=_sha256(xlsx_path),
                extractor="excel", notes=notes,
            )
            s.add(catalog)
            s.flush()

        for cat, model, pos, core in parsed:
            desc_bits = [b for b in (pos, core) if b]
            desc_bits += [f"场景:{sc}" for sc in scenarios_for(model)]
            description = " | ".join(desc_bits)[:2000]

            prod = s.scalar(select(Product).where(Product.model == model[:128]))
            if prod is None:
                prod = Product(model=model[:128])
                s.add(prod)
            prod.name = (pos or model)[:256]
            prod.series = (pos or None)
            prod.category = (cat[:64] or None)
            prod.section = "innovation"          # 名录 = 自主可控承诺型号
            prod.granularity = "model"
            prod.is_domestic = True
            prod.description = description
            s.flush()

            s.add(CatalogEntry(
                catalog_id=catalog.id,
                raw_model_code=model[:128],
                product_id=prod.id,
                match_method="excel",
            ))
            report.models_total += 1
            report.by_category[cat] = report.by_category.get(cat, 0) + 1

    logger.info("名录 %r: imported %d models from %s", name, report.models_total, xlsx_path.name)
    return report


__all__ = ["import_catalog_xlsx", "CatalogXlsxReport"]
