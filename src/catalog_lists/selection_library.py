"""
Import the 全线产品选型库.xlsx as **series-level** products (创新/通用 scope).

Each sheet is `创新产品-<类别>` or `通用产品-<类别>`; each data row is one
*product series* (产品系列). We map every row to a Product tagged
``granularity="series"`` so semantic recall in the 创新/通用 flow returns
series — per the phased plan (创新/通用→系列 now, →型号 in Phase 2).

This is a derivative seed of the SQLite catalog: always safe to re-run
(``catalog import-library``). Rows are upserted by a stable ``model`` key.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

from ..storage import get_db

logger = logging.getLogger(__name__)


@dataclass
class SeriesImportReport:
    source_file: str
    sheets: int = 0
    series_total: int = 0
    skipped_sheets: int = 0
    by_section: dict[str, int] = field(default_factory=dict)
    by_category: dict[str, int] = field(default_factory=dict)


def _parse_sheet_name(name: str) -> tuple[str | None, str]:
    """`创新产品-交换机` -> ("innovation", "交换机"); `通用产品-无线局域网` -> ("general", …)."""
    name = (name or "").strip()
    if name.startswith("创新产品-"):
        return "innovation", name[len("创新产品-"):].strip()
    if name.startswith("通用产品-"):
        return "general", name[len("通用产品-"):].strip()
    return None, name


def _clean(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\r", " ").replace("\n", " ").strip()


def import_selection_library(xlsx_path: str | Path) -> SeriesImportReport:
    """Upsert every series row of the selection-library workbook into the catalog."""
    import openpyxl  # local import: keep openpyxl off the no-Excel path

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    db = get_db()
    report = SeriesImportReport(source_file=str(xlsx_path))
    seen_models: set[str] = set()

    for sheet_name in wb.sheetnames:
        section, category = _parse_sheet_name(sheet_name)
        if section is None:
            report.skipped_sheets += 1
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 3:
            continue  # title + header only, no data
        headers = [_clean(h) for h in rows[1]]
        report.sheets += 1

        for raw in rows[2:]:
            cells = list(raw)
            if len(cells) < 2:
                continue
            series_name = _clean(cells[1])          # col1 = 产品系列/型号/名称
            if not series_name:
                continue                            # blank/spacer row
            positioning = _clean(cells[2]) if len(cells) > 2 else ""  # col2 = 产品定位

            # cols 3.. = category-specific specs -> description + extra_specs
            spec_parts: list[str] = []
            extra: dict[str, str] = {}
            for j in range(3, len(cells)):
                val = _clean(cells[j])
                if not val:
                    continue
                hdr = headers[j] if j < len(headers) and headers[j] else f"col{j}"
                spec_parts.append(f"{hdr}:{val}")
                extra[hdr] = val

            desc_bits = [b for b in (positioning, "；".join(spec_parts)) if b]
            description = " | ".join(desc_bits)[:2000]

            model = series_name[:120]
            if model in seen_models:                # rare cross-sheet name clash
                model = f"{series_name[:108]} ({section})"[:128]
            seen_models.add(model)

            db.upsert_product({
                "model": model,
                "series": series_name[:128],
                "name": series_name[:256],
                "section": section,
                "category": category,
                "granularity": "series",
                "is_domestic": section == "innovation",   # 创新 = 自主可控
                "description": description,
                "extra_specs": extra or None,
            })
            report.series_total += 1
            report.by_section[section] = report.by_section.get(section, 0) + 1
            report.by_category[category] = report.by_category.get(category, 0) + 1

    logger.info("Selection library: imported %d series from %s",
                report.series_total, xlsx_path.name)
    return report


__all__ = ["import_selection_library", "SeriesImportReport"]
